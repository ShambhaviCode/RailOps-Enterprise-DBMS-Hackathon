import csv
import io
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["DB_ENGINE"] = "sqlite"

import app as railops  # noqa: E402
from app import csv_safe  # noqa: E402

MALICIOUS = {
    "train_number": '=HYPERLINK("http://evil.example","Click")',
    "train_name": "+SUM(1,1)",
    "source": "@cmd",
    "destination": "-2+3",
    "capacity": "4242",
}


@pytest.mark.parametrize("value", ["=1+1", "+1", "-1", "@SUM(A1)", "\tx", "\rx"])
def test_csv_safe_neutralises_formula_triggers(value):
    assert csv_safe(value) == "'" + value


@pytest.mark.parametrize("value", ["Mumbai Rajdhani", "12951", "", 1150, None])
def test_csv_safe_leaves_ordinary_values_alone(value):
    assert csv_safe(value) == value


@pytest.fixture
def client_with_malicious_train():
    client = railops.app.test_client()
    client.post("/login", data={"username": "admin", "password": "admin123"})
    client.post("/trains/create", data=MALICIOUS)
    yield client
    with railops.app.app_context():
        for train in railops.Train.query.filter_by(capacity=4242).all():
            railops.db.session.delete(train)
        railops.db.session.commit()


def test_csv_export_does_not_start_cells_with_formulas(client_with_malicious_train):
    body = client_with_malicious_train.get("/reports/export/train/csv").data.decode()
    row = next(r for r in csv.reader(io.StringIO(body)) if r[-1] == "4242")
    assert row[:4] == ["'" + MALICIOUS[k] for k in ("train_number", "train_name", "source", "destination")]


def test_xlsx_export_stores_formulas_as_text(client_with_malicious_train):
    from openpyxl import load_workbook

    data = client_with_malicious_train.get("/reports/export/train/xlsx").data
    sheet = load_workbook(io.BytesIO(data)).active
    row = next(r for r in sheet.iter_rows() if r[-1].value == 4242)
    assert row[0].value == MALICIOUS["train_number"]
    assert all(cell.data_type != "f" for cell in row)
